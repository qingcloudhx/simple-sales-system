from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import re

def extract_price(value):
    """从价格字符串中提取数字"""
    if not value:
        return 0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        match = re.search(r'[\d.]+', value.replace(',', ''))
        if match:
            return float(match.group())
    return 0

def clean_text(value):
    """清理文本，移除多余空格"""
    if not value:
        return ''
    return str(value).strip()

# 读取Excel文件
excel_path = r'C:/Users/xingh/Desktop/武汉赛维尔生物产品目录20260306.xlsx'
wb = load_workbook(excel_path, data_only=True)
ws = wb.active

# 创建新工作簿
new_wb = Workbook()
new_ws = new_wb.active
new_ws.title = "价格表"

# 设置表头样式
header_font = Font(name='Arial', bold=True, size=11)
header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 定义CSV表头
csv_headers = ['中文名称', '英文名称', '销售价', '库存', '分类', '货号', '成本价', '市场价', '图片链接']

# 写入表头
for col_idx, header in enumerate(csv_headers, 1):
    cell = new_ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

# 设置列宽
column_widths = {
    'A': 40,  # 中文名称
    'B': 15,  # 英文名称
    'C': 12,  # 销售价
    'D': 10,  # 库存
    'E': 15,  # 分类
    'F': 15,  # 货号
    'G': 12,  # 成本价
    'H': 12,  # 市场价
    'I': 30,  # 图片链接
}

for col_letter, width in column_widths.items():
    new_ws.column_dimensions[col_letter].width = width

# 数据样式
data_font = Font(name='Arial', size=10)
data_alignment = Alignment(horizontal='center', vertical='center')
data_alignment_left = Alignment(horizontal='left', vertical='center')

# 获取最大行数
max_row = ws.max_row

# 从第3行开始读取（跳过前两行标题）
converted_count = 0
new_row = 2

for row_idx in range(3, max_row + 1):
    分类 = clean_text(ws.cell(row=row_idx, column=1).value)
    项目 = clean_text(ws.cell(row=row_idx, column=2).value)
    货号 = clean_text(ws.cell(row=row_idx, column=3).value)
    名称 = clean_text(ws.cell(row=row_idx, column=4).value)
    品牌 = clean_text(ws.cell(row=row_idx, column=5).value)
    规格 = clean_text(ws.cell(row=row_idx, column=6).value)
    价格 = ws.cell(row=row_idx, column=7).value or ''
    备注 = clean_text(ws.cell(row=row_idx, column=8).value)
    
    # 跳过空行
    if not 名称 and not 货号:
        continue
    
    # 提取价格数字
    original_price = extract_price(价格)
    
    # 生成中文名称: 名称(规格)
    if 规格:
        中文名称 = f"{名称}({规格})"
    else:
        中文名称 = 名称
    
    # 英文名称使用"--"
    英文名称 = '--'
    
    # 价格计算
    销售价 = round(original_price * 0.9, 2) if original_price else 0
    成本价 = round(original_price * 0.7, 2) if original_price else 0
    市场价 = original_price if original_price else 0
    
    # 库存默认1000
    库存 = 1000
    
    # 货号
    货号 = str(货号) if 货号 else ''
    
    # 图片链接为空
    图片链接 = ''
    
    # 写入数据
    row_data = [
        中文名称,
        英文名称,
        销售价,
        库存,
        分类,
        货号,
        成本价,
        市场价,
        图片链接
    ]
    
    for col_idx, value in enumerate(row_data, 1):
        cell = new_ws.cell(row=new_row, column=col_idx, value=value)
        cell.font = data_font
        cell.border = thin_border
        if col_idx == 1:  # 中文名称左对齐
            cell.alignment = data_alignment_left
        else:
            cell.alignment = data_alignment
        
        # 价格列格式化为2位小数
        if col_idx in [3, 7, 8] and value:  # 销售价、成本价、市场价
            cell.number_format = '0.00'
    
    new_row += 1
    converted_count += 1

# 冻结首行
new_ws.freeze_panes = 'A2'

# 保存文件
output_path = r'C:/Users/xingh/Desktop/武汉赛维尔生物产品目录_新价格表.xlsx'
new_wb.save(output_path)

wb.close()
new_wb.close()

print(f"转换完成！")
print(f"共转换 {converted_count} 条数据")
print(f"输出文件: {output_path}")
