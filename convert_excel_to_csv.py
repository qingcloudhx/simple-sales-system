from openpyxl import load_workbook
import re

def extract_price(value):
    """从价格字符串中提取数字"""
    if not value:
        return 0
    # 如果是数字直接返回
    if isinstance(value, (int, float)):
        return float(value)
    # 如果是字符串，尝试提取数字
    if isinstance(value, str):
        # 匹配数字（支持整数和小数）
        match = re.search(r'[\d.]+', value.replace(',', ''))
        if match:
            return float(match.group())
    return 0

# 读取Excel文件
excel_path = r'C:/Users/xingh/Desktop/武汉赛维尔生物产品目录20260306.xlsx'
wb = load_workbook(excel_path, data_only=True)
ws = wb.active

# CSV输出路径
csv_path = r'j:/code/zqz-code/simple-sales-system/files/武汉赛维尔产品目录_converted.csv'

# 定义CSV表头
csv_headers = ['中文名称', '英文名称', '销售价', '库存', '分类', '货号', '成本价', '市场价', '图片链接']

# 打开文件写入
with open(csv_path, 'w', encoding='utf-8-sig') as f:
    # 写入表头
    f.write(','.join(csv_headers) + '\n')
    
    # 获取最大行数
    max_row = ws.max_row
    
    # 从第3行开始读取（跳过前两行标题）
    converted_count = 0
    for row_idx in range(3, max_row + 1):
        # Excel列对应关系:
        # A(1)=分类, B(2)=项目, C(3)=货号, D(4)=名称, E(5)=品牌, F(6)=规格, G(7)=价格, H(8)=备注
        
        分类 = ws.cell(row=row_idx, column=1).value or ''
        项目 = ws.cell(row=row_idx, column=2).value or ''
        货号 = ws.cell(row=row_idx, column=3).value or ''
        名称 = ws.cell(row=row_idx, column=4).value or ''
        品牌 = ws.cell(row=row_idx, column=5).value or ''
        规格 = ws.cell(row=row_idx, column=6).value or ''
        价格 = ws.cell(row=row_idx, column=7).value or ''
        备注 = ws.cell(row=row_idx, column=8).value or ''
        
        # 跳过空行
        if not 名称 and not 货号:
            continue
        
        # 提取价格数字
        original_price = extract_price(价格)
        
        # 转换字段
        中文名称 = 名称
        英文名称 = '--'
        销售价 = str(int(original_price * 0.9)) if original_price else ''  # 销售价为原价九折
        库存 = '1000'  # 默认库存
        分类 = 分类
        货号 = str(货号) if 货号 else ''
        成本价 = str(int(original_price * 0.7)) if original_price else ''  # 成本价为原价七折
        市场价 = str(int(original_price)) if original_price else ''  # 市场价等于原价
        图片链接 = ''
        
        # 组合CSV行
        row_data = [
            中文名称,
            英文名称,
            销售价,
            库存,
            分类,
            货号,
            成本价,
            市场价,
            图片链接
        ]
        
        # 转义CSV特殊字符
        escaped_row = []
        for cell in row_data:
            cell_str = str(cell)
            if ',' in cell_str or '"' in cell_str or '\n' in cell_str:
                cell_str = '"' + cell_str.replace('"', '""') + '"'
            escaped_row.append(cell_str)
        
        f.write(','.join(escaped_row) + '\n')
        converted_count += 1

wb.close()

print(f"转换完成！")
print(f"共转换 {converted_count} 条数据")
print(f"输出文件: {csv_path}")
