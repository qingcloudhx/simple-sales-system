from openpyxl import load_workbook

excel_path = r'C:/Users/xingh/Desktop/武汉赛维尔生物产品目录20260306.xlsx'
wb = load_workbook(excel_path, data_only=True)
ws = wb.active

print("=" * 60)
print("Excel表头 (第一行):")
print("=" * 60)
headers = []
for cell in ws[1]:
    val = cell.value
    headers.append(val)
    print(f"{cell.column}. {val}")

print("\n" + "=" * 60)
print("前3行数据预览:")
print("=" * 60)
for row_idx in range(2, 5):
    row_data = []
    for cell in ws[row_idx]:
        row_data.append(str(cell.value)[:30] if cell.value else "")
    print(f"行{row_idx}: {row_data}")

print("\n" + "=" * 60)
print("CSV目标格式表头:")
print("=" * 60)
csv_headers = ['中文名称', '英文名称', '销售价', '库存', '分类', '货号', '成本价', '市场价', '图片链接']
for i, h in enumerate(csv_headers):
    print(f"{i+1}. {h}")
